"""DPU Filing Assistant: drag in public utility filings, ask questions, and run discovery tasks
(key figures with automatic checks, draft information requests, response tracking, filing
comparison) with every result traced to the page it came from."""
from __future__ import annotations

import hashlib
import io
import os
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st

from rag import answer, discovery, evaluate, ingest, llm, rules, search

st.set_page_config(page_title="DPU Filing Assistant", page_icon="📑", layout="wide")


def secret(name: str, default=None):
    try:
        return st.secrets.get(name, os.environ.get(name, default))
    except Exception:
        return os.environ.get(name, default)


MAX_Q = int(secret("MAX_QUESTIONS", 25))
MAX_MB = 50
ROOT = Path(__file__).parent
SAMPLES = ROOT / "samples"


def md(text: str) -> str:
    """Show dollar signs literally (Streamlit would otherwise read $...$ as math)."""
    return text.replace("\\$", "$").replace("$", "\\$")


def md_stream(gen):
    for piece in gen:
        yield md(piece)


ss = st.session_state
ss.setdefault("docs", {})          # file hash -> Document
ss.setdefault("chat", [])          # {"role", "content", "sources"}
ss.setdefault("asked", 0)          # AI requests used this session
ss.setdefault("unlocked", not secret("APP_PASSWORD"))
ss.setdefault("upload_key", 0)
ss.setdefault("use_samples", False)
ss.setdefault("figs", {})          # (file hash, model) -> list[Figure]
ss.setdefault("irs", None)         # DataFrame of drafted IRs
ss.setdefault("grades", None)      # DataFrame of graded responses
ss.setdefault("eval", None)        # DataFrame of evaluation results

# ------------------------------------------------------------------ access gate
if not ss.unlocked:
    st.title("📑 DPU Filing Assistant")
    pw = st.text_input("Access code", type="password")
    if pw and pw == secret("APP_PASSWORD"):
        ss.unlocked = True
        st.rerun()
    elif pw:
        st.error("Incorrect access code.")
    st.stop()

# ------------------------------------------------------------------ sidebar
with st.sidebar:
    st.header("Settings")
    MODELS = dict(llm.MODELS)
    for prov, override in (("groq", secret("GROQ_MODEL")), ("gemini", secret("GEMINI_MODEL"))):
        if override:
            MODELS = {k: (p_, override if p_ == prov else m) for k, (p_, m) in MODELS.items()}
    owner_keys = {p_: secret(P.key_name) for p_, P in llm.PROVIDERS.items()}
    options = list(MODELS) + [llm.NO_AI]
    default = secret("DEFAULT_MODEL")
    if default not in options:
        default = next((lbl for prov in llm.PREFERENCE if owner_keys[prov]
                        for lbl, (p_, _) in MODELS.items() if p_ == prov), llm.NO_AI)
    model_label = st.selectbox("AI model", options, index=options.index(default),
                               help="Free options run on Groq or Google's free tiers. "
                                    "'No AI' keeps search and the arithmetic checks working without any key.")
    if model_label == llm.NO_AI:
        provider, model, P, key, OWNER_KEY = None, "rules", None, None, False
    else:
        provider, model = MODELS[model_label]
        P = llm.PROVIDERS[provider]
        key = owner_keys[provider]
        OWNER_KEY = bool(key)       # the session limit only applies to the app owner's key
        if not key:
            key = st.text_input(f"{P.name} API key", type="password",
                                help=f"{'Free key' if P.free else 'Create one'} at {P.key_url}. "
                                     "It stays in this browser session.")
            if not key:
                st.info(f"No {P.name} key, so the app runs in no-AI mode: search, figures from tables "
                        "and text patterns, checks, template IRs and comparisons still work.")
    AI = bool(key)
    MODE = model if AI else "rules"     # cache key for extracted figures
    counter = st.empty()
    if st.button("Clear documents and results", width="stretch"):
        for k in ("docs", "figs"):
            ss[k] = {}
        ss.chat, ss.use_samples, ss.irs, ss.grades, ss.eval = [], False, None, None, None
        ss.upload_key += 1          # resets the drag-and-drop box
        st.rerun()
    st.divider()
    st.markdown("**Public filings only.** Uploaded files are processed in memory for this session "
                "and are not stored. " + (P.data_note if AI else "In no-AI mode nothing leaves the app.")
                + " Do not upload confidential or redacted material.")
    st.caption("Created by Cathy Kam")

# ------------------------------------------------------------------ upload
st.title("📑 DPU Filing Assistant")
st.caption("Drag in public filings (GAF/PGAF/OGAF, LDAF, rate cases, information-request responses). "
           "Ask questions, pull key figures with automatic checks, draft information requests, "
           "track responses and compare filings. Everything cites the file and page it came from.")

files = st.file_uploader("Drop filings here", type=[e.lstrip(".") for e in ingest.SUPPORTED],
                         accept_multiple_files=True, key=f"upload_{ss.upload_key}")
current = set()
for f in files or []:
    data = f.getvalue()
    h = hashlib.sha1(data).hexdigest()
    current.add(h)
    if h in ss.docs:
        continue
    if len(data) > MAX_MB * 1e6:
        st.warning(f"{f.name}: larger than {MAX_MB} MB, skipped.")
        continue
    with st.spinner(f"Reading {f.name} …"):
        try:
            ss.docs[h] = ingest.load(f.name, data)
        except Exception as e:
            st.error(f"{f.name}: could not be read ({e}).")

if ss.use_samples:
    for pth in sorted(SAMPLES.glob("sample_*.*")):
        if pth.suffix not in (".pdf", ".xlsx"):
            continue
        data = pth.read_bytes()
        h = hashlib.sha1(data).hexdigest()
        current.add(h)
        if h not in ss.docs:
            ss.docs[h] = ingest.load(pth.name, data)

# files removed from the uploader are dropped from the session
ss.docs = {h: d for h, d in ss.docs.items() if h in current}
docs = list(ss.docs.values())
by_name = {d.name: (h, d) for h, d in ss.docs.items()}

if docs:
    tbl = pd.DataFrame([{
        "File": d.name, "Detected type": d.doc_type,
        "Pages / sheets": d.pages, "Searchable sections": len(d.chunks),
        "IR responses": len({c.ir for c in d.chunks if c.ir}),
        "Dockets": ", ".join(d.dockets[:3]) or "–",
    } for d in docs])
    st.dataframe(tbl, hide_index=True, width="stretch")
    for d in docs:
        if d.warning:
            st.warning(f"**{d.name}:** {d.warning}")
else:
    st.info("Upload one or more filings to start. Supported: PDF, Excel, Word, CSV, text.")
    if st.button("Or try it with three fictional sample filings"):
        ss.use_samples = True
        st.rerun()

# index lives in this browser session only (never shared between visitors)
key_ = tuple(sorted(ss.docs))
if docs and ss.get("index_key") != key_:
    ss.index = search.Index([c for h in key_ for c in ss.docs[h].chunks])
    ss.index_key = key_
index = ss.get("index") if docs else None


# ------------------------------------------------------------------ helpers
def get_client(cost: int = 1):
    """Return an API client, or show why not and return None."""
    if not AI:
        st.error("Choose a model and add its key in the sidebar.")
        return None
    if OWNER_KEY and ss.asked + cost > MAX_Q:
        st.error(f"Request limit for this session reached ({MAX_Q}). Switch to 'No AI' to keep exploring.")
        return None
    return llm.make_client(provider, key)


def run_ai(fn, cost: int = 1):
    """Run an AI task with friendly errors; counts toward the session limit."""
    client = get_client(cost)
    if client is None:
        return None
    try:
        out = fn(client)
    except Exception as e:      # every provider SDK raises its own classes
        st.error(llm.describe_error(e, provider))
        return None
    ss.asked += cost
    return out


def figures_for(names: list[str]) -> list:
    """Extract figures for each named file once per model (or once by rules), reusing results."""
    for n in [n for n in names if (by_name[n][0], MODE) not in ss.figs]:
        h, d = by_name[n]
        if not AI:
            ss.figs[(h, MODE)] = rules.extract_figures(d.chunks)
            continue
        numeric = [c for c in d.chunks if discovery.has_numbers(c)]
        cost = max(1, len(discovery._batches(numeric, P.batch_chars)))
        with st.spinner(f"Extracting figures from {n} …"):
            out = run_ai(lambda client: discovery.extract_figures(
                client, model, d.chunks, max_chars=P.batch_chars, max_tokens=P.max_tokens), cost)
        if out is None:
            break
        ss.figs[(h, MODE)] = out
    return [f for n in names for f in ss.figs.get((by_name[n][0], MODE), [])]


def docx_bytes(title: str, df: pd.DataFrame) -> bytes:
    import docx
    doc = docx.Document()
    doc.add_heading(title, 1)
    for _, r in df.iterrows():
        doc.add_heading(f"Information Request {r['No.']}", 3)
        doc.add_paragraph(str(r["Request"]))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


SEV = {"high": "🔴", "medium": "🟠", "low": "⚪"}

tabs = st.tabs(["💬 Ask", "🔢 Key figures & checks", "📝 Draft IRs", "📬 Track responses",
                "↔️ Compare filings", "🎯 Evaluation"])

# ================================================================== Ask
with tabs[0]:
    scope = st.multiselect("Search within", [d.name for d in docs], default=[d.name for d in docs],
                           help="Limit answers to some of the uploaded files.") if docs else []

    def show_sources(srcs: list, text: str):
        used = answer.cited_ids(text)
        cited = [(n, srcs[n - 1]) for n in used if 0 < n <= len(srcs)]
        if not cited:
            return
        with st.expander(f"Sources cited ({len(cited)})"):
            for n, c in cited:
                st.markdown(f"**[S{n}] {md(c.citation())}**")
                st.text(c.text[:1200] + ("…" if len(c.text) > 1200 else ""))

    for turn in ss.chat:
        with st.chat_message(turn["role"]):
            st.markdown(md(turn["content"]))
            if turn["role"] == "assistant":
                show_sources(turn.get("sources", []), turn["content"])

    SUGGEST = {
        "GAF / cost of gas filing": ["What peak GAF is proposed for residential heating, and how does it "
                                     "compare with last year?", "What drives the change in the cost of gas?",
                                     "What is the reconciliation over- or under-collection?"],
        "Information request response": ["Summarize each information request and the company's answer.",
                                         "Which responses contain dollar figures, and what are they?"],
        "Rate case": ["What revenue increase is requested and what is the proposed return on equity?",
                      "How does the requested increase affect a typical residential heating bill?"],
    }
    pending = None
    if docs and not ss.chat:
        types = {d.doc_type for d in docs}
        sugg = [q for t in types for q in SUGGEST.get(t, [])][:4] or ["Summarize these filings."]
        for col, q in zip(st.columns(len(sugg)), sugg):
            if col.button(q, width="stretch"):
                pending = q

    question = st.chat_input("Ask about the uploaded filings…", disabled=not docs) or pending
    if question:
        with st.chat_message("user"):
            st.markdown(md(question))
        if AI:
            hits = index.search(question, files=scope or None, max_chars=P.context_chars)
        else:
            hits = sorted(index.search(question, k=3, files=scope or None, always_rank=True),
                          key=lambda h: -h[1])[:3]
        with st.chat_message("assistant"):
            if not hits:
                reply, sources = "I couldn't find anything relevant in the selected files.", []
                st.markdown(reply)
            elif not AI:
                parts = ["No AI model is selected, so here are the most relevant passages. "
                         "Choose a model in the sidebar for a written, cited answer."]
                for n, (c, _) in enumerate(hits, start=1):
                    body = c.text if len(c.text) <= 900 else c.text[:900] + "…"
                    quoted = "\n".join("> " + ln for ln in rules.highlight(body, question).splitlines() if ln.strip())
                    parts.append(f"**{n}. {c.citation()}**\n\n{quoted}")
                reply, sources = "\n\n".join(parts), []
                st.markdown(md(reply))
            else:
                def ask(client):
                    stream, srcs = answer.stream_answer(
                        client, model, question, hits,
                        [{"role": t["role"], "content": t["content"]} for t in ss.chat], P.max_tokens)
                    return st.write_stream(md_stream(stream)).replace("\\$", "$"), srcs
                out = run_ai(ask)
                if out is None:
                    st.stop()
                reply, sources = out
                show_sources(sources, reply)
        ss.chat += [{"role": "user", "content": question},
                    {"role": "assistant", "content": reply, "sources": sources}]

    if ss.chat:
        lines = []
        for t in ss.chat:
            lines.append(f"{'Q' if t['role'] == 'user' else 'A'}: {t['content']}\n")
            if t["role"] == "assistant":
                for n in answer.cited_ids(t["content"]):
                    if 0 < n <= len(t.get("sources", [])):
                        lines.append(f"  [S{n}] {t['sources'][n - 1].citation()}\n")
        st.download_button("Download this Q&A", "\n".join(lines).encode(),
                           file_name=f"dpu_filing_qa_{datetime.now():%Y%m%d_%H%M}.txt")

# ================================================================== Key figures
with tabs[1]:
    if AI:
        st.markdown(f"**{model_label}** reads each filing and records every key figure (GAF, LDAF, cost-of-gas "
                    "components, reconciliation, sales, bill impacts) with the exact sentence it came from. "
                    "The app then checks each quote against the page, ties components to totals and flags "
                    "large changes and missing items. The checks are plain arithmetic, not AI.")
    else:
        st.markdown("**No-AI mode:** figures are read from table rows and standard wording "
                    "(\"commodity costs of \\$…\", \"GAF is \\$… per therm\"). This is dependable on tables and "
                    "weaker on free-form prose. The checks are the same arithmetic as with a model.")
    if not docs:
        st.info("Upload filings first.")
    else:
        pick = st.multiselect("Filings", [d.name for d in docs],
                              default=[d.name for d in docs if d.doc_type != "Information request response"]
                              or [d.name for d in docs], key="fig_files")
        done = all((by_name[n][0], MODE) in ss.figs for n in pick)
        if pick and (done or st.button("Extract key figures", type="primary")):
            figs = figures_for(pick)
            if figs:
                issues = discovery.run_checks(figs)
                ss.last_issues = issues
                df = discovery.figures_frame(figs)
                n_ok = int(df["verified"].sum())
                c1, c2, c3 = st.columns(3)
                c1.metric("Figures found", len(df))
                c2.metric("Verified against the page", f"{n_ok}/{len(df)}")
                c3.metric("Issues flagged", len(issues))

                st.subheader("Checks")
                if not issues:
                    st.success("No issues found: components tie to totals and every figure was verified.")
                for i in issues:
                    where = f" — {i.file}, {', '.join(i.locations)}" if i.locations else (f" — {i.file}" if i.file else "")
                    st.markdown(md(f"{SEV[i.severity]} **{i.kind}** {i.description}{where}"))

                st.subheader("Figures")
                show = df.assign(metric=df["metric"].map(discovery.METRICS).fillna(df["metric"]),
                                 verified=df["verified"].map({True: "✓", False: "✗ check"}))
                show = show[["file", "metric", "label", "rate_class", "status", "period", "value", "unit",
                             "verified", "location", "quote"]]
                show.columns = ["File", "Metric", "Label", "Rate class", "Status", "Period", "Value", "Unit",
                                "Verified", "Where", "Quote from filing"]
                st.dataframe(show, hide_index=True, width="stretch",
                             column_config={"Value": st.column_config.NumberColumn(format="%g")})
                st.download_button("Download figures (CSV)", show.to_csv(index=False).encode(),
                                   file_name="key_figures.csv")

# ================================================================== Draft IRs
with tabs[2]:
    st.warning("Draft information requests are a starting point for staff review, built from public "
               "filings for demonstration. Edit before use; issuing discovery follows DPU procedure.")
    if not docs:
        st.info("Upload filings first.")
    else:
        c1, c2, c3, c4 = st.columns([2, 1, 1, 1])
        target = c1.multiselect("Filings to review", [d.name for d in docs],
                                default=[d.name for d in docs if d.doc_type != "Information request response"]
                                or [d.name for d in docs], key="ir_files")
        prefix = c2.text_input("Prefix", "DPU")
        set_no = c3.number_input("Set", 1, 99, 1)
        n_req = c4.number_input("Requests", 3, 25, 10)
        use_checks = st.checkbox("Use automatic checks from the Key figures tab", value=True,
                                 help="Issues found there (tie-outs, large changes, missing items) "
                                      "become the first things asked about.")
        if target and st.button("Draft information requests", type="primary"):
            issues = []
            if use_checks:
                figs = figures_for(target)
                issues = discovery.run_checks(figs) if figs else []
            chunks_t = [c for n in target for c in by_name[n][1].chunks]
            if AI:
                with st.spinner("Drafting …"):
                    out = run_ai(lambda client: discovery.draft_irs(
                        client, model, chunks_t, issues, prefix, int(set_no), int(n_req),
                        max_chars=P.context_chars, max_tokens=P.max_tokens))
            else:
                out = rules.draft_irs(figures_for(target), issues, " ".join(c.text for c in chunks_t),
                                      prefix, int(set_no), int(n_req))
            if out is not None:
                ss.irs = out
        if not AI:
            st.caption("No-AI mode: requests are templates filled from the automatic checks and the figures "
                       "found (reconcile this total, explain this change, provide workpapers for this "
                       "component). A model writes requests tailored to the filing's content.")
        if ss.irs is not None and len(ss.irs):
            st.caption("Edit directly in the table; downloads use your edits.")
            edited = st.data_editor(ss.irs, hide_index=True, width="stretch", num_rows="dynamic",
                                    key="ir_editor",
                                    column_config={"Request": st.column_config.TextColumn(width="large")})
            title = f"Information Requests — Set {int(set_no)}"
            d1, d2, d3 = st.columns(3)
            d1.download_button("Word (.docx)", docx_bytes(title, edited), file_name="information_requests.docx")
            d2.download_button("Text", discovery.irs_to_text(edited, title).encode(),
                               file_name="information_requests.txt")
            d3.download_button("CSV with rationale and sources", edited.to_csv(index=False).encode(),
                               file_name="information_requests.csv")

# ================================================================== Track responses
with tabs[3]:
    st.markdown("Upload the company's responses. Each IR is checked for whether it answers what was asked, "
                "which attachments it relies on (and whether you've uploaded them), and what to follow up on.")
    resp_docs = [d for d in docs if any(c.ir for c in d.chunks)]
    if not resp_docs:
        st.info("No information-request responses detected in the uploaded files.")
    else:
        pick = st.multiselect("Response files", [d.name for d in resp_docs],
                              default=[d.name for d in resp_docs], key="resp_files")
        drafted = list(ss.irs["No."]) if ss.irs is not None and len(ss.irs) else []
        check_drafts = st.checkbox(f"Also flag drafted IRs with no response ({len(drafted)} drafted)",
                                   value=bool(drafted), disabled=not drafted)
        if pick and st.button("Review responses", type="primary"):
            chunks_r = [c for n in pick for c in by_name[n][1].chunks]
            n_ir = len(discovery.ir_groups(chunks_r))
            if AI:
                n_calls = len(discovery.ir_batches(discovery.ir_groups(chunks_r), P.batch_chars))
                with st.spinner(f"Reviewing {n_ir} responses …"):
                    out = run_ai(lambda client: discovery.grade_responses(
                        client, model, chunks_r, [d.name for d in docs], drafted if check_drafts else None,
                        max_chars=P.batch_chars, max_tokens=P.max_tokens), cost=max(1, n_calls))
            else:
                out = rules.grade_responses(chunks_r, [d.name for d in docs], drafted if check_drafts else None)
            if out is not None:
                ss.grades = out
        if not AI:
            st.caption("No-AI mode: responses are checked for objections, deferred answers, attachments "
                       "not uploaded and very short answers. Whether a response fully answers the "
                       "question needs a model (or a person).")
        if ss.grades is not None and len(ss.grades):
            g = ss.grades
            counts = g["Status"].value_counts()
            cols = st.columns(max(1, len(counts)))
            for col, (s, n) in zip(cols, counts.items()):
                col.metric(s.capitalize(), n)
            st.dataframe(g, hide_index=True, width="stretch")
            follow = g[g["Suggested follow-up"].astype(bool)]
            if len(follow):
                st.download_button("Download follow-up requests",
                                   "\n\n".join(f"Follow-up to {r['IR']}: {r['Suggested follow-up']}"
                                               for _, r in follow.iterrows()).encode(),
                                   file_name="follow_up_requests.txt")
            st.download_button("Download tracker (CSV)", g.to_csv(index=False).encode(),
                               file_name="ir_response_tracker.csv")

# ================================================================== Compare
with tabs[4]:
    st.markdown("Pick last year's filing and this year's. Figures are extracted from both and matched by "
                "metric and rate class, so every change is shown with the page each number came from.")
    names = [d.name for d in docs]
    if len(names) < 2:
        st.info("Upload at least two filings (for example, last year's and this year's GAF filing).")
    else:
        c1, c2 = st.columns(2)
        prior = c1.selectbox("Prior filing", names, index=0)
        cur = c2.selectbox("Current filing", names, index=1)
        if prior == cur:
            st.warning("Choose two different filings.")
        elif st.button("Compare", type="primary") or all((by_name[n][0], MODE) in ss.figs for n in (prior, cur)):
            figs = figures_for([prior, cur])
            if figs:
                cmp = discovery.compare(figs, prior, cur)
                st.dataframe(cmp, hide_index=True, width="stretch", column_config={
                    "value (prior)": st.column_config.NumberColumn("Prior", format="%g"),
                    "value (current)": st.column_config.NumberColumn("Current", format="%g"),
                    "Change": st.column_config.NumberColumn(format="%+g"),
                    "Change %": st.column_config.NumberColumn(format="%+.1f%%")})
                st.download_button("Download comparison (CSV)", cmp.to_csv(index=False).encode(),
                                   file_name="filing_comparison.csv")

# ================================================================== Evaluation
with tabs[5]:
    st.markdown("How often is the assistant right? Load a question set with known answers "
                "(a CSV: `id, question, files, expected_values, expected_location`), then score "
                "retrieval (free) or full answers and citations, optionally against a whole-document "
                "baseline, which is what pasting the filings into a chat app would give.")
    sample_q = ROOT / "eval" / "questions_sample.csv"
    up = st.file_uploader("Question set (CSV)", type=["csv"], key="qset")
    try:
        qs = evaluate.load_questions(up.getvalue() if up else sample_q)
    except Exception as e:
        qs = []
        st.error(f"Could not read the question set: {e}")
    if not up:
        st.caption(f"Using the sample set ({len(qs)} questions on the fictional sample filings). "
                   "Load the samples above to run it.")
    have = {d.name for d in docs}
    runnable = [q for q in qs if set(q.files) <= have]
    st.write(f"{len(runnable)} of {len(qs)} questions have all their files uploaded.")
    c1, c2, c3 = st.columns(3)
    k = c1.number_input("Passages per question (k)", 2, 30, 8)
    base = c2.checkbox("Include whole-document baseline", value=True)
    if AI and P.free:
        st.caption(f"{P.name}'s free tier has a per-minute token limit, so scoring runs slowly and the "
                   f"baseline is capped at {min(evaluate.BASELINE_CHARS, P.context_chars * 2):,} characters "
                   "of each filing.")
    st.download_button("Template for scoring ChatGPT / Claude.ai by hand",
                       evaluate.manual_template(qs).encode(), file_name="manual_comparison.csv")
    if runnable:
        b1, b2 = st.columns(2)
        all_chunks = [c for d in docs for c in d.chunks]
        if b1.button("Score retrieval (free)"):
            ss.eval = evaluate.results_frame(evaluate.run(runnable, all_chunks, k=int(k)))
        cost = len(runnable) * (2 if base else 1)
        if b2.button(f"Score answers ({cost} API requests)", type="primary", disabled=not AI,
                     help="Needs a model. Uses your own API key; on the shared demo key this may exceed "
                          "the session limit."):
            bar = st.progress(0.0)
            res = run_ai(lambda client: evaluate.run(
                runnable, all_chunks, client=client, model=model, k=int(k), baseline=base,
                max_chars=P.context_chars, max_tokens=P.max_tokens,
                baseline_chars=min(evaluate.BASELINE_CHARS, P.context_chars * 2),
                progress=lambda i, n, q: bar.progress(i / n, text=f"{q.id}: {q.question[:60]}")), cost)
            if res is not None:
                ss.eval = evaluate.results_frame(res)
    if ss.eval is not None:
        st.dataframe(evaluate.summary(ss.eval), hide_index=True)
        st.dataframe(ss.eval, hide_index=True, width="stretch")
        st.download_button("Download report (Markdown)",
                           evaluate.report_markdown(ss.eval, "Evaluation", f"Model {model_label}, k={k}").encode(),
                           file_name="evaluation_report.md")
        st.download_button("Download results (CSV)", ss.eval.to_csv(index=False).encode(),
                           file_name="evaluation_results.csv")

counter.caption(f"AI requests this session: {ss.asked}" + (f" / {MAX_Q}" if OWNER_KEY else ""))
