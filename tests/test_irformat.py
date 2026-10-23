"""Information requests in the Department's issued format.

Reference: Second Set of Information Requests of the Department of Public Utilities to the Local
Distribution Companies, D.P.U. 26-PGAF Dockets (October 2, 2026), a public document."""
import io
import re

import pandas as pd

from rag import compliance as C
from rag import discovery as D
from rag import ingest
from rag import irformat as F

PUBLISHED_2_1 = (
    "Refer to the Company’s September 17, 2026 Peak GAF filing, the Company’s August 3, 2026 LDAF filing, any "
    "revisions to the Company’s GAF or LDAF filings, and the Company’s response to information request DPU 1-1. "
    "For each illustrative residential LDAF that the Company calculated to reallocate 30, 40, and 50 percent of "
    "costs, please provide: (a) updated residential bill impacts to also include the impact due to the proposed "
    "peak GAF; and (b) an updated residential “all-in” bill impact analysis, comparing all proposed peak rates "
    "(e.g., GAFs, illustrative LDAFs, RDAFs, GSEAFs, GSERAFs, base distribution rates due to a performance-based "
    "rate adjustment or a base distribution rate case) for effect November 1, 2026, to those rates that were in "
    "effect on November 1, 2025.")


def flat(t):
    return re.sub(r"\s+", " ", t).strip()


def test_rebuilds_a_published_request_exactly():
    text = F.compose(
        "the Company's September 17, 2026 Peak GAF filing, the Company's August 3, 2026 LDAF filing, any revisions "
        "to the Company's GAF or LDAF filings, and the Company's response to information request DPU 1-1",
        "for each illustrative residential LDAF that the Company calculated to reallocate 30, 40, and 50% of costs, "
        "please provide",
        ["Updated residential bill impacts to also include the impact due to the proposed peak GAF.",
         'an updated residential "all-in" bill impact analysis, comparing all proposed peak rates (e.g. GAFs, '
         "illustrative LDAFs, RDAFs, GSEAFs, GSERAFs, base distribution rates due to a performance-based rate "
         "adjustment or a base distribution rate case) for effect November 1, 2026, to those rates that were in "
         "effect on November 1, 2025"])
    assert flat(text) == PUBLISHED_2_1
    assert "DPU 1-1.  For each" in text                       # two spaces after a sentence, as issued
    assert F.lint(text) == []


def test_style_rules():
    t = F.style_text('The Company\'s "all-in" rate rose 7.9% (e.g. R-3). See D.P.U. 26-PGAF at 2. Done.')
    assert t == "The Company’s “all-in” rate rose 7.9 percent (e.g., R-3).  See D.P.U. 26-PGAF at 2.  Done."
    body, subs = F.split_request("Refer to x at 2.  Please provide: (a) one; (b) two; and (c) three.")
    assert body == "Refer to x at 2.  Please provide:" and subs == ["one;", "two; and", "three."]
    assert F.format_request("Refer to x at 2. Please provide:\n(a) One\n(b) Two") == \
        "Refer to x at 2.  Please provide:\n(a) one; and\n(b) two."
    assert F.label("DPU", 2, 3) == "DPU 2-3" and F.label("DPU-LDC", 2, 1) == "DPU-LDC 2-1"
    assert F.normalize_id("DPU 2-1") == F.normalize_id("dpu-2–1") == "DPU-2-1"
    assert F.IRSet(set_no=2).ordinal == "SECOND"


def test_lint_points():
    assert any("Refer to" in n for n in F.lint("Please provide workpapers."))
    assert any("pinpoint" in n for n in F.lint("Refer to the filing.  Please provide workpapers."))
    assert any("Please" in n for n in F.lint("Refer to the filing at 2.  Workpapers are needed."))
    assert any("single subpart" in n for n in F.lint("Refer to the filing at 2.  Please provide:\n(a) one."))


def test_pinpoints_and_defaults():
    assert F.pinpoint("p. 2 (table)") == "at 2"
    assert F.pinpoint("Sheet 'Schedule 3', rows 5-10") == "at Schedule 3, rows 5-10"
    assert F.guess_company(["SAMPLE DOCUMENT – fictional\nSample Gas Company\nD.P.U. 99"]) == "Sample Gas Company"
    assert F.guess_docket(["25-GREC-06", "26-109", "26-PGAF-NSTAR"]) == "26-PGAF-NSTAR"
    assert F.default_reference("x", "Local Distribution Adjustment Clause", "GAF / cost of gas filing",
                               "August 2026") == "the Company’s August 2026 LDAF filing"


META = F.IRSet("NSTAR Gas Company", "26-PGAF-NSTAR", 2, "DPU", "October 2, 2026")


def table():
    return pd.DataFrame([
        {"No.": "", "Topic": "Bill impacts", "Request": PUBLISHED_2_1.replace(": (a)", ":\n(a)")
            .replace("; and (b)", "; and\n(b)"), "Why": "", "Based on": ""},
        {"No.": "", "Topic": "Cash-out", "Request": "Refer to the Company's GAF filing, Section VI. Please explain "
         "how the Company recovers the variance.", "Why": "", "Based on": ""}])


def test_text_layout():
    t = F.to_text(table(), META)
    assert "SECOND SET OF INFORMATION REQUESTS OF THE" in t and "TO NSTAR GAS COMPANY" in t
    assert "Pursuant to 220 CMR 1.06(5)(c), the Department of Public Utilities (“Department”) submits to" in t
    assert re.search(r"^DPU 2-1 +Refer to", t, re.M) and re.search(r"^ {14}\(a\) +updated residential", t, re.M)
    assert t.rstrip().endswith("Dated:  October 2, 2026") and t.startswith("DRAFT")


def test_word_layout_matches_issued_sets():
    import docx
    from docx.shared import Inches
    d = docx.Document(io.BytesIO(F.to_docx(table(), META)))
    paras = [p for p in d.paragraphs if p.text.strip()]
    assert d.styles["Normal"].font.name == "Times New Roman" and d.styles["Normal"].font.size.pt == 12
    assert paras[0].text.splitlines()[0] == "SECOND SET OF INFORMATION REQUESTS OF THE"
    assert paras[0].alignment == 1                                           # centered
    req_head = next(p for p in paras if p.text == "Requests")
    assert req_head.runs[0].bold and req_head.runs[0].underline
    first = next(p for p in paras if p.text.startswith("DPU 2-1\t"))
    assert first.paragraph_format.left_indent == Inches(1.0)
    assert first.paragraph_format.first_line_indent == Inches(-1.0)
    sub = next(p for p in paras if p.text.startswith("(a)\t"))
    assert sub.paragraph_format.left_indent == Inches(1.5) and sub.paragraph_format.first_line_indent == Inches(-0.5)
    assert paras[-1].text == "Dated:  October 2, 2026"
    sec = d.sections[0]
    assert sec.different_first_page_header_footer
    assert "D.P.U. 26-PGAF-NSTAR\tPage" in sec.header.paragraphs[0].text
    assert 'w:instr="PAGE"' in sec.header.paragraphs[0]._p.xml
    assert "DRAFT" in sec.first_page_header.paragraphs[0].text and "DRAFT" in sec.footer.paragraphs[0].text


def test_ai_drafts_are_composed_in_format():
    from types import SimpleNamespace as NS

    class Fake:
        def __init__(self):
            self.messages, self.calls = self, []

        def create(self, **kw):
            self.calls.append(kw)
            return NS(content=[NS(type="tool_use", name="record_requests", input={"requests": [
                {"topic": "Bill impacts", "refer_to": "the Company's Peak GAF filing at Schedule 2",
                 "request": "please provide", "subparts": ["Bill impacts for R-1", "bill impacts for R-3."],
                 "rationale": "x", "source_ids": [1]},
                {"topic": "One", "refer_to": "the Company's Peak GAF filing at 2", "request": "Please provide",
                 "subparts": ["workpapers for the LNG cost"], "rationale": "x", "source_ids": []}]})])

    chunks = ingest.load("sample_pgaf_filing.pdf", open(
        __import__("pathlib").Path(__file__).resolve().parents[1] / "samples" / "sample_pgaf_filing.pdf", "rb").read()).chunks
    fake = Fake()
    df = D.draft_irs(fake, "m", chunks, [], "DPU", 2, refs={"sample_pgaf_filing.pdf": "the Company’s Peak GAF filing"})
    assert df["Request"].iloc[0] == ("Refer to the Company’s Peak GAF filing at Schedule 2.  Please provide:\n"
                                     "(a) bill impacts for R-1; and\n(b) bill impacts for R-3.")
    assert df["Request"].iloc[1] == "Refer to the Company’s Peak GAF filing at 2.  Please provide workpapers for the LNG cost."
    assert "cite as “the Company’s Peak GAF filing”" in fake.calls[0]["messages"][0]["content"]
    assert "percent" in fake.calls[0]["system"] and "Refer to the Company’s September 17, 2026" in fake.calls[0]["system"]


def test_tracking_matches_department_numbering_to_responses():
    chunks = ingest.load("r.txt", b"Information Request DPU-1-2\nRequest: x\nResponse: y is provided here in full "
                                  b"detail with the numbers requested.").chunks
    g = D.add_unanswered([], D.ir_groups(chunks), ["DPU 1-2", "DPU 1-3"])
    assert list(g["IR"]) == ["DPU 1-3"]                     # DPU 1-2 was answered as "DPU-1-2"


def test_compliance_follow_ups_in_format():
    root = __import__("pathlib").Path(__file__).resolve().parents[1] / "samples"
    load = lambda n: ingest.load(n, (root / n).read_bytes())
    orders = [load("order_sample_98-PGAF-SAMPLE.pdf"), load("order_sample_cover_letter.pdf")]
    filing = [load(n) for n in ("sample_pgaf_filing.pdf", "sample_ir_responses.pdf", "sample_schedule3.xlsx")]
    refs = {d.name: C.order_reference(d) for d in orders}
    assert refs == {"order_sample_98-PGAF-SAMPLE.pdf": "the Department’s Order in D.P.U. 98-PGAF-SAMPLE",
                    "order_sample_cover_letter.pdf": "the Company’s filing cover letter"}
    irs = C.follow_up_irs(C.results_frame(C.check(orders, filing)), "DPU", 2, 1, refs, "the Company’s Peak GAF filing")
    first = irs["Request"].iloc[0]
    assert first.startswith("Refer to the Department’s Order in D.P.U. 98-PGAF-SAMPLE at 1, which states: “ORDERED:")
    assert "the filing shows 0.8302" in first and F.lint(first) == []
    assert any("which lists “Schedule 4 – Hedging Program Summary” as enclosed.  Please provide" in r
               for r in irs["Request"])


def test_ldac_templates_cite_other_dockets_not_the_filing_itself():
    from rag import rules
    from rag import verify as V
    root = __import__("pathlib").Path(__file__).resolve().parents[1] / "samples"
    d = ingest.load("ldac.pdf", (root / "real_NSTAR_Gas_LDAC_DPU_26-PGAF-NSTAR_corrected.pdf").read_bytes())
    own = F.guess_docket(d.dockets, [V.page_texts(d)[0][1]])
    assert own == "26-PGAF-NSTAR"                              # the header's docket, not 26-OGAF-NSTAR
    figs = V.schedule_figures(d)
    df = rules.draft_irs(figs, D.run_checks(figs), " ".join(c.text for c in d.chunks), "DPU", 1, 10,
                         {d.name: "the Company’s August 2026 LDAF filing"},
                         [(d.name, loc, s) for loc, s in V.document_schedules(d)], own)
    lines = [re.search(r"Line (\d+)", r).group(1) for r in df["Request"]]
    assert lines == ["7", "9", "10", "18", "20", "21"]         # lines citing other dockets only
    assert not any("hedg" in r.lower() for r in df["Request"])  # no cost-of-gas templates on an LDAC
    assert "cites “D.P.U. 26-99 filed July 30, 2026, effective November 1, 2026” for Line 7, Pension/PBOP Factor" \
        in df["Request"].iloc[0]
    assert all(F.lint(r) == [] for r in df["Request"])


def test_filing_date_ignores_dates_in_docket_citations():
    root = __import__("pathlib").Path(__file__).resolve().parents[1] / "samples"
    d = ingest.load("ldac.pdf", (root / "real_NSTAR_Gas_LDAC_DPU_26-PGAF-NSTAR_corrected.pdf").read_bytes())
    assert C.filing_date([d])[0] == "August 2026"              # "Filed: August 2026", not Line 7's citation
    assert F.default_reference(d.name, " ".join(t for _, t in __import__("rag.verify", fromlist=["x"]).page_texts(d)),
                               d.doc_type, "August 2026") == "the Company’s August 2026 LDAF filing"
