"""Turn uploaded filings into searchable chunks with exact citations.

Each chunk remembers its file and location (PDF page, Excel sheet and rows,
Word paragraph range) so every answer can point back to the source.

Regulatory filings need a few special touches:
* Information-request (IR) responses are split per question, so
  "DPU-1-3" retrieves exactly that question and answer.
* PDF tables are extracted as rows (pipe-separated) so dollar figures and
  their labels stay together.
* Excel schedules keep the header row with every block of rows.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

MAX_CHARS = 1800        # target chunk size
OVERLAP = 250           # characters carried into the next chunk
ROWS_PER_CHUNK = 40     # Excel/CSV rows per chunk

# DPU-1-3, AG-2-14, DPU 3-1, IR DPU-1-3 ...
IR_PATTERN = re.compile(
    r"\b(?:Information\s+Request\s+)?((?:DPU|AG|DOER|LEI|CLF|NCLC|[A-Z]{2,5})[-\s]?\d{1,2}[-–]\d{1,3})\b")
DOCKET_PATTERN = re.compile(r"\bD\.?\s?P\.?\s?U\.?\s*(\d{2}-[A-Z0-9-]+)", re.I)

DOC_TYPES = [
    ("Information request response", r"information\s+request|response\s+to\s+(dpu|ag)|\bIR\b\s*[A-Z]+-\d"),
    ("GAF / cost of gas filing", r"gas\s+adjustment\s+factor|\bp?gaf\b|\bogaf\b|cost\s+of\s+gas|\bldaf\b"),
    ("Rate case", r"revenue\s+requirement|cost\s+of\s+service|rate\s+base|base\s+distribution\s+rate|"
                  r"pre-?filed\s+(direct\s+)?testimony"),
]


@dataclass
class Chunk:
    text: str
    file: str
    location: str              # "p. 12", "Sheet 'Sch 1', rows 2-41", "¶ 10-25"
    page: int | None = None
    ir: str | None = None      # e.g. "DPU-1-3"
    kind: str = "text"         # text | table | spreadsheet
    meta: dict = field(default_factory=dict)

    def citation(self) -> str:
        bits = [self.file, self.location]
        if self.ir:
            bits.append(f"IR {self.ir}")
        return " · ".join(bits)


@dataclass
class Document:
    name: str
    doc_type: str
    pages: int
    chunks: list[Chunk]
    dockets: list[str]
    warning: str | None = None


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split(text: str, max_chars: int = MAX_CHARS, overlap: int = OVERLAP) -> list[str]:
    """Split on paragraph boundaries, then sentences, with overlap."""
    if len(text) <= max_chars:
        return [text] if text.strip() else []
    paras = re.split(r"\n\s*\n", text)
    out, cur = [], ""
    for p in paras:
        if len(p) > max_chars:      # very long paragraph: split on sentences
            parts = re.split(r"(?<=[.;:])\s+", p)
        else:
            parts = [p]
        for part in parts:
            if len(cur) + len(part) + 2 > max_chars and cur:
                out.append(cur.strip())
                cur = cur[-overlap:] + "\n" + part
            else:
                cur = f"{cur}\n\n{part}" if cur else part
    if cur.strip():
        out.append(cur.strip())
    return out


def classify(text: str, name: str) -> str:
    sample = (name + " " + text[:20000]).lower()
    for label, pat in DOC_TYPES:
        if re.search(pat, sample, re.I):
            return label
    return "Other filing"


def find_dockets(text: str) -> list[str]:
    return sorted({m.group(1).upper() for m in DOCKET_PATTERN.finditer(text[:50000])})


def split_by_ir(text: str) -> list[tuple[str | None, str]]:
    """Split a response document at each IR heading ("Information Request DPU-1-3")."""
    heads = [m for m in re.finditer(
        r"(?im)^\s*(?:information\s+request|request\s+no\.?|ir)?\s*[:#]?\s*"
        r"((?:DPU|AG|DOER|[A-Z]{2,5})[-\s]?\d{1,2}[-–]\d{1,3})\b", text)]
    if len(heads) < 2:
        m = IR_PATTERN.search(text)
        return [(m.group(1).replace(" ", "-") if m else None, text)]
    parts = []
    if heads[0].start() > 0 and text[:heads[0].start()].strip():
        parts.append((None, text[:heads[0].start()]))
    for i, h in enumerate(heads):
        end = heads[i + 1].start() if i + 1 < len(heads) else len(text)
        parts.append((h.group(1).replace(" ", "-").replace("–", "-").upper(), text[h.start():end]))
    return parts


# --------------------------------------------------------------------------
# readers
# --------------------------------------------------------------------------
def read_pdf(name: str, data: bytes) -> Document:
    import pdfplumber

    chunks, all_text, pages, empty = [], [], 0, 0
    current_ir = None
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            text = _clean(text)
            if not text:
                empty += 1
            all_text.append(text)
            for ir, part in split_by_ir(text) if text else []:
                current_ir = ir or current_ir
                for piece in _split(part):
                    chunks.append(Chunk(piece, name, f"p. {i}", page=i, ir=current_ir))
            try:
                for t in page.extract_tables() or []:
                    rows = [" | ".join(str(c or "").strip() for c in r) for r in t if any(r)]
                    if len(rows) >= 2:
                        body = "\n".join(rows)
                        for piece in _split(body):
                            chunks.append(Chunk(piece, name, f"p. {i} (table)", page=i,
                                                ir=current_ir, kind="table"))
            except Exception:
                pass
    joined = "\n".join(all_text)
    warning = None
    if pages and empty / pages > 0.5:
        warning = (f"{empty} of {pages} pages have no extractable text (likely scanned images). "
                   "Run OCR first (e.g. Adobe 'Recognize Text') for those pages to be searchable.")
    return Document(name, classify(joined, name), pages, chunks, find_dockets(joined), warning)


def _frame_chunks(df: pd.DataFrame, name: str, where: str, first_row: int = 2) -> list[Chunk]:
    df = df.dropna(how="all").dropna(axis=1, how="all")
    if df.empty:
        return []
    df = df.astype(str).replace({"nan": "", "None": "", "NaT": ""})
    header = " | ".join(str(c) for c in df.columns)
    out = []
    for start in range(0, len(df), ROWS_PER_CHUNK):
        block = df.iloc[start:start + ROWS_PER_CHUNK]
        lines = [" | ".join(v.strip() for v in row) for row in block.values.tolist()]
        text = f"{where}\nColumns: {header}\n" + "\n".join(lines)
        out.append(Chunk(text[:MAX_CHARS * 2], name,
                         f"{where.split(' (')[0]}, rows {first_row + start}-"
                         f"{first_row + start + len(block) - 1}",
                         kind="spreadsheet"))
    return out


def read_excel(name: str, data: bytes) -> Document:
    xls = pd.ExcelFile(io.BytesIO(data))
    chunks, text = [], []
    for sheet in xls.sheet_names:
        raw = pd.read_excel(xls, sheet, header=None)
        if raw.dropna(how="all").empty:
            continue
        # use the first mostly-filled row as the header (schedules often have titles above it)
        filled = raw.notna().sum(axis=1)
        hdr = int(filled[filled >= max(2, filled.max() * 0.6)].index[0])
        title = " ".join(str(v) for v in raw.iloc[:hdr].stack().dropna().tolist()
                         if str(v).strip() and str(v) != "nan")[:200]
        df = raw.iloc[hdr + 1:].copy()
        df.columns = [str(c) if pd.notna(c) else f"col{j}" for j, c in enumerate(raw.iloc[hdr])]
        where = f"Sheet '{sheet}'" + (f" ({title})" if title else "")
        chunks += _frame_chunks(df, name, where, first_row=hdr + 2)
        text.append(where + " " + " ".join(map(str, raw.iloc[:20].stack().dropna().tolist())))
    joined = "\n".join(text)
    return Document(name, classify(joined, name), len(xls.sheet_names), chunks, find_dockets(joined))


def read_csv(name: str, data: bytes) -> Document:
    df = pd.read_csv(io.BytesIO(data))
    chunks = _frame_chunks(df, name, "CSV")
    return Document(name, classify(" ".join(df.columns), name), 1, chunks, [])


def read_docx(name: str, data: bytes) -> Document:
    import docx

    d = docx.Document(io.BytesIO(data))
    paras = [p.text for p in d.paragraphs]
    chunks, buf, start = [], [], 0
    for i, p in enumerate(paras + ["\n"]):
        buf.append(p)
        if sum(len(x) for x in buf) > MAX_CHARS or i == len(paras):
            text = _clean("\n".join(buf))
            for ir, part in split_by_ir(text) if text else []:
                for piece in _split(part):
                    chunks.append(Chunk(piece, name, f"¶ {start + 1}-{i + 1}", ir=ir))
            buf, start = [], i + 1
    for ti, t in enumerate(d.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        for piece in _split("\n".join(rows)):
            chunks.append(Chunk(piece, name, f"table {ti}", kind="table"))
    joined = "\n".join(paras)
    return Document(name, classify(joined, name), len(paras), chunks, find_dockets(joined))


def read_text(name: str, data: bytes) -> Document:
    text = _clean(data.decode("utf-8", errors="ignore"))
    chunks = [Chunk(p, name, f"part {i}", ir=ir)
              for i, (ir, part) in enumerate(split_by_ir(text), start=1) for p in _split(part)]
    return Document(name, classify(text, name), 1, chunks, find_dockets(text))


READERS = {".pdf": read_pdf, ".xlsx": read_excel, ".xlsm": read_excel, ".xls": read_excel,
           ".csv": read_csv, ".docx": read_docx, ".txt": read_text, ".md": read_text}
SUPPORTED = sorted(READERS)


def load(name: str, data: bytes) -> Document:
    ext = Path(name).suffix.lower()
    if ext not in READERS:
        raise ValueError(f"Unsupported file type {ext}. Supported: {', '.join(SUPPORTED)}")
    doc = READERS[ext](name, data)
    if not doc.chunks and not doc.warning:
        doc.warning = "No text could be extracted from this file."
    return doc
